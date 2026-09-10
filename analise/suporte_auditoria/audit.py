from pathlib import Path
from collections import Counter, defaultdict
import openpyxl, json, re, unicodedata, hashlib

SOURCE=Path('/Users/joaomvalente/Downloads/Campanha_EA_2026_REV-006.xlsx')
OUT=Path(__file__).parent
OUT.mkdir(parents=True, exist_ok=True)
w=openpyxl.load_workbook(SOURCE, data_only=False)
def norm(x):
    if x is None: return ''
    s=str(x).strip().upper()
    s=''.join(c for c in unicodedata.normalize('NFD',s) if not unicodedata.combining(c))
    return re.sub(r'\s+',' ',s)
def ref(s,r): return f"'{s}'!{r}"
def rows_of(s):
    for row in s.iter_rows(min_row=3):
        if any(norm(c.value) for c in row[:8]): yield row

groups=defaultdict(list)
records=[]
for s in w:
    typ='municipio' if norm(s['A2'].value)=='ARTICULADOR' else 'recorte' if norm(s['A2'].value)=='CIDADE' and norm(s['B2'].value)=='ARTICULADOR' else 'indice'
    groups[typ].append(s.title)
    if typ=='indice': continue
    for row in rows_of(s):
        r=row[0].row
        if typ=='municipio':
            values=[s.title]+[s.cell(r,c).value for c in range(1,9)]
            coords=[None]+[s.cell(r,c).coordinate for c in range(1,9)]
            if norm(s['E2'].value)=='CONTATO' and norm(s['F2'].value)=='REGIAO':
                values[5],values[6]=values[6],values[5]
                coords[5],coords[6]=coords[6],coords[5]
        else:
            values=[s.cell(r,c).value for c in range(1,8)]+[s['B1'].value or s.title,s.cell(r,8).value]
            coords=[s.cell(r,c).coordinate for c in range(1,8)]+['B1',s.cell(r,8).coordinate]
        records.append(dict(sheet=s.title,row=r,type=typ,values=values,coords=coords,norm=[norm(x) for x in values]))

data={'sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'groups':groups,'sheet_counts':{s.title:sum(r['sheet']==s.title for r in records) for s in w},'count_types':dict(Counter(r['type'] for r in records))}
data['structure']=[{'sheet':s.title,'max_row':s.max_row,'max_col':s.max_column,'header':[s.cell(2,c).value for c in range(1,s.max_column+1)],'merged':[str(x) for x in s.merged_cells.ranges],'formulas':[c.coordinate for row in s for c in row if c.data_type=='f'],'validations':len(s.data_validations.dataValidation),'tables':len(s.tables),'hidden':s.sheet_state,'filter':s.auto_filter.ref} for s in w]
data['missing']={t:{str(i):sum(not r['norm'][i] for r in records if r['type']==t) for i in range(9)} for t in ('municipio','recorte')}
data['empty_strings']=[dict(sheet=s.title,cell=c.coordinate) for s in w for row in s for c in row if isinstance(c.value,str) and not norm(c.value)]
data['whitespace']=[dict(sheet=s.title,cell=c.coordinate) for s in w for row in s for c in row if isinstance(c.value,str) and c.value!=c.value.strip() and norm(c.value)]
for scope in ('municipio','recorte','both'):
    bucket=defaultdict(list)
    for r in records:
        if scope!='both' and r['type']!=scope:continue
        key=tuple(r['norm'])
        bucket[key].append(ref(r['sheet'],r['row']))
    data[f'duplicates_full_{scope}']=[v for v in bucket.values() if len(v)>1]

# Comparison between different views. The party-contact label is excluded to
# avoid claiming aliases are automatically the same person or organization.
shared_fields=[0,1,2,3,4,5,6,8]
buckets=defaultdict(list)
for r in records: buckets[tuple(r['norm'][i] for i in shared_fields)].append(r)
overlap=[v for v in buckets.values() if {r['type'] for r in v}=={'municipio','recorte'}]
data['shared_fields_overlap']={'groups':len(overlap),'municipio_rows':sum(r['type']=='municipio' for v in overlap for r in v),'recorte_rows':sum(r['type']=='recorte' for v in overlap for r in v),'examples':[[ref(r['sheet'],r['row']) for r in v] for v in overlap[:15]]}
data['unmatched_shared_fields']={t:[ref(r['sheet'],r['row']) for v in buckets.values() if {r['type'] for r in v}=={t} for r in v] for t in ('municipio','recorte')}

phone_stats={}
for t in ('municipio','recorte'):
    phone_stats[t]={}
    for i in (3,6):
        vals=[r for r in records if r['type']==t and r['norm'][i]]
        invalid=[]; nonstring=[]; repeated=defaultdict(list)
        for r in vals:
            val=r['values'][i]
            digits=re.sub(r'\D','',str(val))
            if len(digits) in (12,13) and digits.startswith('55'):digits=digits[2:]
            # Syntax only; no validation of subscriber ownership or service.
            if len(digits) not in (10,11) or re.search(r'[A-Za-z]',str(val)):
                invalid.append(dict(ref=ref(r['sheet'],r['coords'][i]),value=val,digits=len(digits)))
            if not isinstance(val,str): nonstring.append(dict(ref=ref(r['sheet'],r['coords'][i]),value=val,type=type(val).__name__))
            repeated[digits].append(dict(ref=ref(r['sheet'],r['coords'][i]),name=r['values'][2 if i==3 else 4]))
        phone_stats[t][str(i)]={'filled':len(vals),'invalid':invalid,'nonstring':nonstring,'repeat_groups':[v for k,v in repeated.items() if k and len(v)>1]}
data['phone_stats']=phone_stats

links=[]
for s in w:
    for row in s:
        for c in row:
            if c.hyperlink:
                h=c.hyperlink
                target=h.location or h.target
                if not target:continue
                internal=target.startswith('#') or h.location is not None
                if internal:
                    loc=target.lstrip('#'); sheet=loc.rsplit('!',1)[0].strip("'").replace("''", "'") if '!' in loc else None
                    links.append(dict(source=ref(s.title,c.coordinate),label=c.value,target=target,target_sheet=sheet,exists=sheet in w.sheetnames,internal=True))
                else: links.append(dict(source=ref(s.title,c.coordinate),label=c.value,target=target,internal=False))
data['links']=links
data['broken_links']=[l for l in links if l['internal'] and not l['exists']]
data['indices']={s.title:[dict(row=r,cells=[s.cell(r,c).value for c in range(1,s.max_column+1)]) for r in range(1,s.max_row+1) if any(norm(s.cell(r,c).value) for c in range(1,s.max_column+1))] for s in w if s.title in groups['indice']}
data['formula_errors']=[dict(sheet=s.title,cell=c.coordinate,value=c.value) for s in w for row in s for c in row if c.data_type=='e' or (c.data_type=='f' and '#REF!' in c.value)]
data['records_without_leader']=[dict(sheet=r['sheet'],row=r['row'],values=r['values']) for r in records if not r['norm'][4]]

OUT.joinpath('audit.json').write_text(json.dumps(data,ensure_ascii=False,indent=2,default=str))
OUT.joinpath('records.json').write_text(json.dumps(records,ensure_ascii=False,indent=2,default=str))
print(json.dumps({k:data[k] for k in ('sha256','groups','count_types','missing','shared_fields_overlap','broken_links','records_without_leader')},ensure_ascii=False,indent=2,default=str))
