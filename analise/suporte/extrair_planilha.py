from pathlib import Path
import openpyxl, json, collections, hashlib

source = Path('/Users/joaomvalente/Downloads/Campanha_EA_2026_REV-006.xlsx')
out = Path(__file__).parent
out.mkdir(parents=True, exist_ok=True)
w = openpyxl.load_workbook(source, data_only=False)
result = {'source': source.name, 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'sheets': []}
lines = []
for s in w:
    cells = []
    lines.append('\n## ' + s.title)
    for row in s:
        populated = []
        for c in row:
            if c.value is not None:
                v = c.value
                item = {'cell': c.coordinate, 'value': v, 'type': c.data_type, 'format': c.number_format}
                if c.hyperlink:
                    item['hyperlink'] = {'target':c.hyperlink.target,'location':c.hyperlink.location}
                if c.comment:
                    item['comment'] = c.comment.text
                cells.append(item)
                populated.append(c.coordinate+'='+str(v).replace('\n',' / '))
        if populated:
            lines.append(' | '.join(populated))
    result['sheets'].append({'name':s.title,'dimensions':s.calculate_dimension(),'state':s.sheet_state,
        'cells':cells,'merged':[str(x) for x in s.merged_cells.ranges],
        'validations':[{'range':str(x.sqref),'type':x.type,'formula1':x.formula1,'formula2':x.formula2} for x in s.data_validations.dataValidation],
        'hidden_rows':[k for k,v in s.row_dimensions.items() if v.hidden],
        'hidden_columns':[k for k,v in s.column_dimensions.items() if v.hidden],
        'tables':list(s.tables),'charts':len(s._charts),'images':len(s._images),
        'conditional_formats':len(s.conditional_formatting),'protected':s.protection.sheet})
(out/'planilha_extraida.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str))
(out/'planilha_celulas.txt').write_text('\n'.join(lines))
print(json.dumps({'sheets':len(w.sheetnames),'cells':sum(len(s['cells']) for s in result['sheets']),
    'formulas':sum(c['type']=='f' for s in result['sheets'] for c in s['cells']),
    'images':sum(s['images'] for s in result['sheets']), 'charts':sum(s['charts'] for s in result['sheets']),
    'comments':sum('comment' in c for s in result['sheets'] for c in s['cells']),
    'hyperlinks':sum('hyperlink' in c for s in result['sheets'] for c in s['cells']),
    'defined_names':list(w.defined_names),'external_links':len(w._external_links),'sha256':result['sha256']},ensure_ascii=False,indent=2))
